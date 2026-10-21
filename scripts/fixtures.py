"""
Sample fixtures for the smoke test.

Provides a small but realistic resume and job description, plus a builder for
the DOCX version of the resume.
"""

from pathlib import Path

import docx

RESUME_LINES = [
    "Jane Doe",
    "jane.doe@example.com",
    "Skills",
    "Python, SQL, AWS, Docker, PostgreSQL",
    "Experience",
    "Data Engineer at Acme Corp",
    "- Built ETL pipelines on AWS with Python and PostgreSQL",
    "- Improved processing speed by 30% and cut infrastructure costs by 15%",
]
RESUME = "\n".join(RESUME_LINES)

JOB = "\n".join((
    "We are hiring a data engineer with Python, SQL, AWS and Airflow.",
    "Docker, Kubernetes and PostgreSQL required, plus good communication.",
))


def write_docx(directory: Path) -> Path:
    """Write the sample resume as a DOCX file with a paragraph and a table."""
    document = docx.Document()
    document.add_paragraph(RESUME)
    table = document.add_table(rows=1, cols=2)
    for index, value in enumerate(("Phone", "+1 555 0100")):
        table.cell(0, index).text = value
    path = directory / "resume.docx"
    document.save(path)
    return path
