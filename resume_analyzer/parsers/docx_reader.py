"""
Plain-text extraction from DOCX resumes.

Resumes are frequently built with tables, so table cell text is collected in
addition to the body paragraphs.
"""

import io

import docx


def extract_text_from_docx(data: bytes) -> str:
    """Return paragraph text plus table cell text found in a DOCX document."""
    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
